import openpyxl
import json
import os

FILE = "data/Penawaran Harga UJE.xlsx"

wb = openpyxl.load_workbook(FILE, data_only=True)

result = {}

for ws in wb.worksheets:
    rows = []

    for row in ws.iter_rows():
        cells = []

        for cell in row:
            if cell.value is not None:
                cells.append({
                    "cell": cell.coordinate,
                    "value": cell.value
                })

        if cells:
            rows.append(cells)

    result[ws.title] = {
        "max_row": ws.max_row,
        "max_column": ws.max_column,
        "rows": rows
    }

output_file = "data/extracted_excel.json"

with open(output_file, "w", encoding="utf-8") as f:
    json.dump(result, f, ensure_ascii=False, indent=2, default=str)

print(f"Berhasil membaca: {FILE}")
print(f"Hasil disimpan ke: {output_file}")